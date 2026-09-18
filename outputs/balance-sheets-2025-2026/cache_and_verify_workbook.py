"""Evaluate this workbook's small formula vocabulary and populate preview caches."""
import json
import re
import zipfile
from decimal import Decimal as D
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

root=Path(__file__).resolve().parent
path=root/'balance-sheets-2025-2026-review.xlsx'
wb=load_workbook(path,data_only=False)
memo={}
active=set()

def cells(sheet,area):
    left,top,right,bottom=range_boundaries(area)
    return [cell for row in wb[sheet].iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right) for cell in row]

def number(value):
    if value is None: return D(0)
    return D(str(value))

def value(sheet,coordinate):
    key=(sheet,coordinate)
    if key in memo: return memo[key]
    assert key not in active, f'Circular formula {key}'
    active.add(key)
    cell=wb[sheet][coordinate]
    result=expression(sheet,cell.value[1:]) if cell.data_type=='f' else cell.value
    active.remove(key)
    memo[key]=result
    return result

def expression(sheet,text):
    text=text.strip()
    match=re.fullmatch(r"SUMIF\('([^']+)'!([A-Z]+\d+:[A-Z]+\d+),\"([^\"]+)\",'([^']+)'!([A-Z]+\d+:[A-Z]+\d+)\)-SUMIF\('([^']+)'!([A-Z]+\d+:[A-Z]+\d+),\"([^\"]+)\",'([^']+)'!([A-Z]+\d+:[A-Z]+\d+)\)",text)
    if match:
        args=match.groups()
        def sumif(a):
            criteria_sheet,criteria_area,criterion,sum_sheet,sum_area=a
            pairs=zip(cells(criteria_sheet,criteria_area),cells(sum_sheet,sum_area),strict=True)
            return sum((number(value(sum_sheet,s.coordinate)) for c,s in pairs if value(criteria_sheet,c.coordinate)==criterion),D(0))
        return sumif(args[:5])-sumif(args[5:])
    match=re.fullmatch(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',text)
    if match: return sum((number(value(sheet,c.coordinate)) for c in cells(sheet,match[1])),D(0))
    match=re.fullmatch(r"'([^']+)'!([A-Z]+\d+)",text)
    if match: return number(value(match[1],match[2]))
    match=re.fullmatch(r'([A-Z]+\d+)([+-])([A-Z]+\d+)',text)
    if match:
        left=number(value(sheet,match[1]));right=number(value(sheet,match[3]))
        return left+right if match[2]=='+' else left-right
    match=re.fullmatch(r'-([A-Z]+\d+)',text)
    if match: return -number(value(sheet,match[1]))
    if re.fullmatch(r'[A-Z]+\d+',text): return number(value(sheet,text))
    if re.fullmatch(r'-?\d+(\.\d+)?',text): return D(text)
    raise ValueError(f'Unsupported formula: {sheet} {text}')

cached={}
for sheet in wb:
    for row in sheet:
        for cell in row:
            if cell.data_type=='f': cached[(sheet.title,cell.coordinate)]=number(value(sheet.title,cell.coordinate)).quantize(D('.01'))
totals=json.loads((root/'summary.json').read_text(encoding='utf-8'))['totals']
for sheet_name,date in [('المركز 2025','2025-12-31'),('المركز 2026 أغسطس','2026-08-31')]:
    required={'إجمالي الأصول':'assets','إجمالي الالتزامات':'liabilities','حقوق الملكية شاملة النتيجة غير المقفلة':'total_equity','إجمالي الالتزامات وحقوق الملكية':'liabilities_equity','فرق المعادلة':'difference'}
    for row in wb[sheet_name]:
        if row[0].value in required:
            assert cached[(sheet_name,row[2].coordinate)]==D(totals[date][required[row[0].value]])
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
temporary=path.with_suffix('.cached.xlsx')
with zipfile.ZipFile(path) as src,zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        data=src.read(item.filename)
        m=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',item.filename)
        if m:
            sheet_name=wb.sheetnames[int(m[1])-1]
            tree=ET.fromstring(data)
            for cell in tree.findall('.//s:c',ns):
                key=(sheet_name,cell.attrib['r'])
                if key not in cached: continue
                v=cell.find('s:v',ns)
                if v is None: v=ET.SubElement(cell,'{'+ns['s']+'}v')
                v.text=str(cached[key])
            data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
        dst.writestr(item,data)
temporary.replace(path)
check=load_workbook(path,data_only=True)
for (sheet,coordinate),expected in cached.items():
    assert D(str(check[sheet][coordinate].value))==expected
print(f'Validated and cached {len(cached)} formulas; both statements match independent totals.')
