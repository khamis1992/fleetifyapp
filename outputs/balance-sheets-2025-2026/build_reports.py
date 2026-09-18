import json
import sys
import hashlib
from decimal import Decimal
from pathlib import Path
from html import escape
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
sys.stdout.reconfigure(encoding='utf-8')
D = Decimal
source_bytes = (ROOT / 'source-evidence.json').read_bytes()
source = json.loads(source_bytes, parse_float=D)
source_hash = hashlib.sha256(source_bytes).hexdigest()
company = source['company']['name_ar'].strip()
dates = ['2025-12-31', '2026-08-31']
audit = json.loads((ROOT / 'audit-evidence.json').read_text(encoding='utf-8'))
counts = {dt: {status: next(r['entries'] for r in audit['per_cutoff'] if r['as_of']==dt and r['status']==status) for status in ('posted','draft')} for dt in dates}
translated = {'20201': 'دفعات مقدمة من العملاء'}
type_names = {'assets':'الأصول', 'liabilities':'الالتزامات', 'equity':'حقوق الملكية', 'revenue':'الإيرادات', 'expenses':'المصروفات'}

def name(r):
    return (r.get('account_name_ar') or translated.get(r['account_code']) or r['account_name']).strip()

def amount(r):
    net = D(r['debits']) - D(r['credits'])
    return -net if r['account_type'] in ('liabilities', 'equity', 'revenue') else net

def money(value):
    value = D(value).quantize(D('.01'))
    return f'({abs(value):,.2f})' if value < 0 else f'{value:,.2f}'

def period(dt):
    rows = [r for r in source['balances'] if r['as_of'] == dt]
    totals = {kind: sum((amount(r) for r in rows if r['account_type'] == kind), D(0)) for kind in type_names}
    totals['residual'] = totals['revenue'] - totals['expenses']
    totals['total_equity'] = totals['equity'] + totals['residual']
    totals['liabilities_equity'] = totals['liabilities'] + totals['total_equity']
    totals['difference'] = totals['assets'] - totals['liabilities_equity']
    assert totals['difference'] == 0, (dt, totals)
    assert sum(D(r['debits']) - D(r['credits']) for r in rows) == 0
    control = next(r for r in audit['per_cutoff'] if r['as_of']==dt and r['status']=='posted')
    assert sum(D(r['debits']) for r in rows) == D(control['line_debit'])
    assert sum(D(r['credits']) for r in rows) == D(control['line_credit'])
    assert sum(r['line_count'] for r in rows) == int(control['lines'])
    return rows, totals

def review_notes(dt):
    rows, t = period(dt)
    contra = next(r for r in rows if r['account_code'] == '11211')
    notes = [
        'لا تظهر أرصدة مرحلة للسيارات والأصول الثابتة أو مجمع الإهلاك أو رأس المال. ومن سجل الأسطول الحالي، تفتقد 221 مركبة من أصل 223 تاريخ الشراء وتكلفته، مع أصلين ثابتين فقط بقيمة إجمالية 170,000 ر.ق. هذا فحص اكتمال للسجل الحالي، وليس حصرًا تاريخيًا لملكية المركبات؛ يلزم مستندات وسجل أصول معتمد.',
        f'حساب 11211 «إيجار منتهي بالتملك» مصنف أصلًا، ورصيده دائن بمبلغ {money(abs(amount(contra)))} ر.ق. أُبقي بإشارته السالبة داخل الأصول؛ يلزم تفسيره ومراجعة تصنيفه قبل الاعتماد.',
        f'استُبعدت {counts[dt]["draft"]:,} مسودة قيد مؤرخة حتى تاريخ القائمة. يجب تحديد ما يخص الفترة منها وما يلزم تصحيحه أو ترحيله بعد المراجعة.',
        'توجد حركات على حسابات من المستوى الثاني، وتفتقر معظم الحسابات ذات الحركة إلى تصنيف فرعي. يلزم مراجعة حسابات الترحيل وتصنيف المتداول وغير المتداول. لا توجد ضمن الحركات المستخرجة مصروفات رواتب أو إيجار أو تأمين أو وقود أو إهلاك؛ غيابها لا يثبت عدم تحمل الشركة لها.',
        'لم تُطابق الأرصدة مع كشوف البنوك والصندوق أو كشوف العملاء والموردين، ولم يُتحقق من اكتمال المصروفات والرواتب والقروض والالتزامات. توازن الأرقام وحده لا يثبت اكتمالها أو صحتها المستندية.',
    ]
    if dt == '2026-08-31':
        notes.append('توجد قيود إقفال مرحلة خلال 2026؛ مبلغ الأرباح المحتجزة هنا من حساب حقوق الملكية كما هو مسجل. بند النتيجة غير المقفلة هو صافي الأرصدة المتبقية في حسابات الإيرادات والمصروفات، وليس تقرير أرباح سنة 2026 مستقلًا. يلزم مراجعة الإقفال وترحيل الأرصدة.')
    else:
        notes.append('لا يظهر رصيد مرحّل في حسابات حقوق الملكية حتى نهاية 2025. تظهر حقوق الملكية في هذه المسودة من صافي أرصدة الإيرادات والمصروفات غير المقفلة؛ وهذا لا يثبت رأس المال الفعلي أو صحة تسويات الإقفال.')
    reversal_count = 2480 if dt == '2025-12-31' else 2168
    notes.append(f'توجد {reversal_count:,} قيود أصلية مشمولة عُكست بقيود مؤرخة بعد تاريخ القائمة. أُبقي الأصل حتى تاريخ عكسه وفق التاريخ المحاسبي المسجل؛ يلزم تقييم التصحيحات اللاحقة قبل اعتماد قائمة معدلة عن الفترة.')
    return notes

CSS = '''
@page { size: A4; margin: 14mm 15mm 14mm; }
* { box-sizing:border-box; }
body { font-family:Arial, sans-serif; font-size:14px; line-height:1.48; color:#17202a; margin:0; direction:rtl; }
.page { break-after:page; } .page:last-child { break-after:auto; }
.company { font-size:19px; font-weight:700; color:#000; }
.cr { font-size:12px; color:#4b5563; margin-top:3px; }
h1 { font-size:25px; color:#000; margin:17px 0 3px; }
h2 { font-size:21px; color:#000; margin:16px 0 5px; }
.period { font-size:16px; margin:0 0 6px; font-weight:bold; }
.status { font-weight:bold; color:#854d0e; margin:7px 0; }
.lead { margin:7px 0 12px; font-size:13px; }
table { border-collapse:collapse; width:100%; margin:10px 0; font-size:13px; }
th,td { border:1px solid #d9d9d9; padding:5px 9px; vertical-align:middle; }
th { background:#233c54; color:white; font-weight:bold; text-align:right; }
.section td { background:#eaf0f5; font-weight:bold; }
.total td { background:#f1f3f5; font-weight:bold; }
.grand td { font-weight:bold; border-top:2px solid #233c54; }
.num { direction:ltr; text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; width:31%; }
.code { direction:ltr; font-size:11px; text-align:center; width:19%; overflow-wrap:anywhere; }
.small { font-size:11.5px; color:#4b5563; }
.notes { margin:10px 0; padding-right:21px; }
.notes li { margin-bottom:8px; }
.meta { font-size:12px; color:#4b5563; margin-top:13px; padding-top:8px; border-top:1px solid #d9d9d9; }
.hash { font-family:Consolas, monospace; direction:ltr; word-break:break-all; font-size:9px; }
a { color:#233c54; }
'''

def tr(label, value, code='', cls=''):
    return f'<tr class="{cls}"><td>{escape(label)}</td><td class="code">{escape(code)}</td><td class="num">{money(value)}</td></tr>'

def header():
    return f'<div class="company">{escape(company)}</div><div class="cr">سجل تجاري {escape(source["company"]["commercial_register"])} · دولة قطر · المبالغ بالريال القطري</div>'

def make_html(dt):
    rows, t = period(dt)
    title_date = datetime.fromisoformat(dt).strftime('%d / %m / %Y')
    body = ''
    for kind in ('assets', 'liabilities'):
        body += f'<tr class="section"><td colspan="3">{type_names[kind]}</td></tr>'
        for r in rows:
            if r['account_type'] == kind and amount(r) != 0:
                body += tr(name(r), amount(r), r['account_code'])
        body += tr('إجمالي ' + type_names[kind], t[kind], cls='total')
    body += '<tr class="section"><td colspan="3">حقوق الملكية وفق الأرصدة المسجلة</td></tr>'
    equity_rows = [r for r in rows if r['account_type']=='equity' and amount(r) != 0]
    if equity_rows:
        body += ''.join(tr(name(r), amount(r), r['account_code']) for r in equity_rows)
    else:
        body += tr('رصيد حسابات حقوق الملكية المرحلة', D(0))
    body += tr('صافي أرصدة الإيرادات والمصروفات غير المقفلة', t['residual'])
    body += tr('إجمالي حقوق الملكية', t['total_equity'], cls='total')
    body += tr('إجمالي الالتزامات وحقوق الملكية', t['liabilities_equity'], cls='grand')
    body += tr('فرق المعادلة الحسابية', t['difference'], cls='total')
    notes = ''.join(f'<li>{escape(n)}</li>' for n in review_notes(dt))
    controls = ''.join(tr(label, value) for label,value in [
        ('إجمالي الأصول وفق التصنيف المسجل', t['assets']),
        ('إجمالي الالتزامات وحقوق الملكية', t['liabilities_equity']),
        ('الفرق', t['difference']),
    ])
    html = f'''<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>مسودة قائمة المركز المالي {dt}</title><style>{CSS}</style></head><body>
<section class="page">{header()}<h1>مسودة قائمة المركز المالي</h1><p class="period">كما في {title_date}</p>
<p class="status">للمراجعة المحاسبية والقانونية · غير مدققة وغير معتمدة</p>
<p class="lead">تعرض هذه المسودة أرصدة القيود المرحلة المسجلة بالنظام حتى التاريخ أعلاه. توجد ملاحظات جوهرية بشأن اكتمال الأصول ورأس المال وتصنيف الأرصدة، موضحة في الصفحة الثانية.</p>
<table><thead><tr><th>البند</th><th class="code">رمز الحساب</th><th class="num">المبلغ ر.ق</th></tr></thead><tbody>{body}</tbody></table>
<p class="small">المبالغ بين قوسين سالبة. صافي النتيجة غير المقفلة محسوب من أرصدة الإيرادات ناقص المصروفات، ولا يمثل بالضرورة ربح السنة وحدها. لم يُعرض رصيد دائن مصنف أصلًا باعتباره أصلًا موجبًا.</p>
<div class="meta">المصدر سجلات النظام · استخراج بتاريخ 18 / 09 / 2026 بتوقيت قطر · الملاحظات في الصفحة الثانية جزء من هذه المسودة</div>
</section>
<section class="page">{header()}<h2>ملاحظات المراجعة وحدود الاستخراج</h2><p class="period">القائمة كما في {title_date}</p><p class="status">لا تمثل هذه المسودة اعتمادًا محاسبيًا أو تقرير تدقيق</p>
<ol class="notes">{notes}</ol>
<h2>نطاق الاستخراج والتحقق</h2>
<p>شمل الاستخراج <b>{counts[dt]['posted']:,}</b> قيدًا مرحّلًا حتى تاريخ القائمة، بما فيها قيود الإقفال والعكس، دون إضافة الرصيد المخزن مرة أخرى. لم يظهر فرق مدين ودائن أو عدم تطابق بين مجموع البنود وإجمالي القيد.</p>
<p class="small">استُبعدت المسودات والملغاة والقيود اللاحقة. الأرقام مستخرجة من السجلات الحالية وقد تتضمن إدخالات لاحقة بتاريخ محاسبي سابق. لا يعني غياب رصيد من النظام أنه صفر في الواقع.</p>
<p class="small">تقتصر الحزمة على مسودتي مركز مالي وملف أرصدة. المجموعة الكاملة للقوائم وفق <a href="https://www.ifrs.org/issued-standards/list-of-standards/ias-1-presentation-of-financial-statements/">IAS 1</a> تشمل قوائم وإيضاحات أخرى؛ ولا يُدّعى أن هذه المسودة تستوفيها.</p>
<div class="meta">مرجع النسخة BS-{dt}-DRAFT · تاريخ الاستخراج 18 / 09 / 2026 بتوقيت قطر<br>بصمة ملف الأرصدة لأغراض المطابقة فقط وليست توقيع اعتماد<div class="hash">{source_hash}</div></div>
</section></body></html>'''
    (ROOT / f'balance-sheet-{dt}-draft.html').write_text(html,encoding='utf-8')
    return t

wb=Workbook()
ws=wb.active
ws.title='دليل الملف'
ws.append(['مسودتا قائمة المركز المالي لشركة العراف'])
ws.append(['التواريخ','2025-12-31 و2026-08-31'])
ws.append(['الحالة','للمراجعة فقط؛ غير مدققة وغير معتمدة'])
ws.append(['العملة','QAR — ريال قطري'])
ws.append(['المصدر','القيود المرحلة لكل تاريخ؛ دون جمع الرصيد المخزن مرة ثانية'])
ws.append(['النتيجة غير المقفلة','رصيد الإيرادات ناقص رصيد المصروفات، وليس بالضرورة ربح السنة'])
ws.append(['العلامات','الرصيد العادي سالب عند وجود رصيد عكسي؛ المبالغ السالبة بين قوسين'])
ws.append(['اكتمال البيانات','لا تظهر أرصدة مرحلة للسيارات والإهلاك ورأس المال؛ يلزم استكمال المراجعة'])
ws.append(['تاريخ الاستخراج UTC',source['extracted_at']])
ws.append(['SHA256',source_hash])
ws.append(['تنبيه','توازن الأرقام لا يثبت اكتمال الميزانية؛ راجع ورقة الملاحظات'])
summaries={}
for dt in dates:
    rows,t=period(dt)
    summaries[dt]={k:str(v) for k,v in t.items()}
    make_html(dt)
    label='2025' if dt.startswith('2025') else '2026 أغسطس'
    raw=wb.create_sheet('أرصدة '+label)
    raw.append(['رمز الحساب','اسم الحساب','النوع','إجمالي المدين','إجمالي الدائن','صافي مدين ناقص دائن','الرصيد حسب طبيعة النوع','مستوى الحساب','عدد البنود'])
    for i,r in enumerate(rows,2):
        is_credit=r['account_type'] in ('liabilities','equity','revenue')
        raw.append([r['account_code'],name(r),type_names[r['account_type']],float(r['debits']),float(r['credits']),f'=D{i}-E{i}',f'=-F{i}' if is_credit else f'=F{i}',r['account_level'],r['line_count']])
    n=raw.max_row
    raw.append(['','إجمالي الحركة','',f'=SUM(D2:D{n})',f'=SUM(E2:E{n})',f'=SUM(F2:F{n})','','',f'=SUM(I2:I{n})'])
    raw.auto_filter.ref=f'A1:I{n}'
    raw.freeze_panes='D2'
    report=wb.create_sheet('المركز '+label)
    report.append(['مسودة قائمة المركز المالي',dt])
    report.append(['غير مدققة وغير معتمدة','QAR'])
    report.append(['البند','رمز الحساب','المبلغ'])
    ranges={}
    for kind in ('assets','liabilities','equity'):
        report.append([type_names[kind],'',''])
        start=report.max_row+1
        for raw_i,r in enumerate(rows,2):
            if r['account_type']==kind:
                report.append([name(r),r['account_code'],f"='أرصدة {label}'!G{raw_i}"])
        if report.max_row < start:
            report.append(['لا يظهر رصيد مرحل','','=0'])
        end=report.max_row
        report.append(['إجمالي '+type_names[kind],'',f'=SUM(C{start}:C{end})'])
        ranges[kind]=report.max_row
    report.append(['صافي الإيرادات والمصروفات غير المقفلة','',f'=SUMIF(\'أرصدة {label}\'!C2:C{n},"الإيرادات",\'أرصدة {label}\'!G2:G{n})-SUMIF(\'أرصدة {label}\'!C2:C{n},"المصروفات",\'أرصدة {label}\'!G2:G{n})'])
    residual_row=report.max_row
    report.append(['حقوق الملكية شاملة النتيجة غير المقفلة','',f'=C{ranges["equity"]}+C{residual_row}'])
    total_equity_row=report.max_row
    report.append(['إجمالي الالتزامات وحقوق الملكية','',f'=C{ranges["liabilities"]}+C{total_equity_row}'])
    report.append(['فرق المعادلة','',f'=C{ranges["assets"]}-C{report.max_row}'])
    report.freeze_panes='C4'
notes=wb.create_sheet('ملاحظات المراجعة')
notes.append(['التاريخ','الملاحظة'])
for dt in dates:
    for note in review_notes(dt): notes.append([dt,note])
notes.append(['المنهج','قائمة 2026 حتى أغسطس وليست نهاية السنة. لا يُدّعى اكتمال القوائم أو مطابقتها لجميع متطلبات IFRS.'])
snapshot=wb.create_sheet('ملخص القيم عند الاستخراج')
snapshot.append(['البند','2025-12-31','2026-08-31'])
for key,title in [('assets','إجمالي الأصول'),('liabilities','إجمالي الالتزامات'),('equity','حسابات حقوق الملكية'),('residual','النتيجة غير المقفلة'),('total_equity','حقوق الملكية شاملة النتيجة'),('liabilities_equity','الالتزامات وحقوق الملكية'),('difference','فرق المعادلة')]:
    snapshot.append([title]+[float(D(summaries[dt][key])) for dt in dates])
snapshot.append(['قيم ثابتة وقت الاستخراج؛ الحسابات التفصيلية في الأوراق الأخرى'])
border=Border(*[Side(style='thin',color='D9D9D9')]*4)
for sheet in wb:
    sheet.sheet_view.rightToLeft=True
    sheet.sheet_view.showGridLines=False
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape' if sheet.max_column>4 else 'portrait'
    sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1
    sheet.page_setup.fitToHeight=0
    sheet.print_title_rows='1:1'
    for row in sheet:
        for cell in row:
            cell.font=Font(name='Arial',size=11,color='17202A')
            cell.alignment=Alignment(horizontal='right',vertical='center',wrap_text=True)
            cell.border=border
            if isinstance(cell.value,(int,float)) or cell.data_type=='f': cell.number_format='#,##0.00;(#,##0.00);0.00'
        sheet.row_dimensions[row[0].row].height=27
    for cell in sheet[1]:
        cell.fill=PatternFill('solid',fgColor='233C54')
        cell.font=Font(name='Arial',size=12,color='FFFFFF',bold=True)
    for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=23
    sheet.column_dimensions['A'].width=42
    if sheet.title.startswith('أرصدة'):
        sheet.column_dimensions['A'].width=25
        sheet.column_dimensions['B'].width=41
        sheet.column_dimensions['H'].width=15
        sheet.column_dimensions['I'].width=15
        for row in sheet.iter_rows(min_row=2):
            for col in (7,8): row[col].number_format='0'
    if sheet.title.startswith('المركز'):
        sheet.column_dimensions['A'].width=59
        sheet.column_dimensions['B'].width=25
        sheet.column_dimensions['C'].width=25
        for row in sheet.iter_rows(min_row=4):
            if str(row[0].value).startswith(('إجمالي','فرق','حقوق الملكية شاملة')):
                for cell in row:
                    cell.font=Font(name='Arial',size=11,bold=True,color='17202A')
                    cell.fill=PatternFill('solid',fgColor='EAF0F5')
    if sheet.title in ('دليل الملف','ملاحظات المراجعة'):
        sheet.column_dimensions['A'].width=31
        sheet.column_dimensions['B'].width=112
        for i in range(2,sheet.max_row+1): sheet.row_dimensions[i].height=58 if sheet.title=='ملاحظات المراجعة' else 36
        sheet.freeze_panes='B2'
wb.calculation.fullCalcOnLoad=True
wb.calculation.forceFullCalc=True
wb.save(ROOT/'balance-sheets-2025-2026-review.xlsx')
check=load_workbook(ROOT/'balance-sheets-2025-2026-review.xlsx',data_only=False)
assert len(check.sheetnames)==7
assert sum(c.data_type=='f' for s in check for row in s for c in row)>70
(ROOT/'summary.json').write_text(json.dumps({'totals':summaries,'source_sha256':source_hash,'status':'draft_unreviewed'},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'totals':summaries,'sheets':check.sheetnames},ensure_ascii=False))
