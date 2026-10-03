import docx, os, subprocess
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def update_alattiya_letter():
    doc = docx.Document('/home/ubuntu/upload/قالبالعراف.docx')
    
    for s in doc.sections:
        s.top_margin = Cm(2.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)

    def set_cell_background(cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D7E1E8"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr.append(borders)

    def set_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2):
        p.alignment = align
        pPr = p._p.get_or_add_pPr()
        bidi = pPr.find(qn('w:bidi'))
        if bidi is None:
            bidi = OxmlElement('w:bidi')
            pPr.append(bidi)
        bidi.set(qn('w:val'), '1')
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15

    def add_p(text, size=10.5, bold=False, color=(23, 50, 71), align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2):
        p = doc.add_paragraph()
        set_rtl(p, align=align, space_after=space_after)
        r = p.add_run(text)
        r.font.name = "Amiri"
        rPr = r._element.get_or_add_rPr()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rPr.rFonts.set(qn(attr), "Amiri")
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = RGBColor(*color)
        return p

    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)

    # PAGE 1: OFFICIAL LETTER
    p_meta = doc.add_paragraph()
    set_rtl(p_meta, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    r1 = p_meta.add_run("التاريخ: 04 / 10 / 2026 م          الرقم الإشاري: ALARAF/ATT/2026-10")
    r1.font.name = "Amiri"
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(100, 116, 139)

    add_p("السادة / شركة العطية للسيارات والتجارة المحترمين", size=12, bold=True, color=(41, 75, 98), space_after=1)
    add_p("وكلاء سيارات شانجان (Changan) — دولة قطر", size=10, bold=True, color=(100, 116, 139), space_after=2)
    add_p("تحية طيبة وبعد،،،", size=10.5, bold=True, space_after=4)

    # Subject Box
    table_subj = doc.add_table(rows=1, cols=1)
    table_subj.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_s = table_subj.cell(0, 0)
    set_cell_background(cell_s, "E1ECF2")
    set_cell_margins(cell_s, top=80, bottom=80, left=120, right=120)
    p_s = cell_s.paragraphs[0]
    set_rtl(p_s, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    r_s = p_s.add_run("الموضوع: طلب عدم ممانعة نقل ملكية سيارات شانجان وتحويل اللوحات من (ليموزين) إلى (خصوصي)")
    r_s.font.name = "Amiri"
    r_s.font.size = Pt(10.5)
    r_s.bold = True
    r_s.font.color.rgb = RGBColor(41, 75, 98)

    add_p("", space_after=2)

    add_p("بالإشارة إلى الموضوع أعلاه، وإلى أسطول سيارات شانجان (Changan Alsvin 2024) البالغ عددها (20) مركبة، المسجلة باسم شركتنا / شركة العراف لتأجير السيارات (ذ.م.م) والممولة لديكم / تحت إشرافكم؛", size=10, space_after=3)
    
    add_p("نرجو التكرم بالموافقة وإصدار كتابكم المعتمد بعدم الممانعة والموجه إلى إدارة المرور والدوريات (وزارة الداخلية) لاتخاذ إجراءات نقل الملكية وتعديل فئة اللوحات المرورية وفق التالي:", size=10, space_after=3)

    add_p("1. نقل الملكية من:", size=10, bold=True, color=(41, 75, 98), space_after=1)
    add_p("   - شركة العراف لتأجير السيارات (ذ.م.م) — سجل تجاري: (146832) — قيد منشأة: (17-2015-86).", size=9.5, space_after=2)

    add_p("2. نقل الملكية إلى:", size=10, bold=True, color=(41, 75, 98), space_after=1)
    add_p("   - شركة جبريل لتوزيع الأغذية والمشروبات (ذ.م.م) — سجل تجاري: (246953) — قيد منشأة: (17-3190-85).", size=9.5, space_after=2)

    add_p("3. تعديل نوع التسجيل واللوحات المرورية:", size=10, bold=True, color=(41, 75, 98), space_after=1)
    add_p("   - تحويل اللوحات المرورية من فئة (ليموزين / نقل ركاب) إلى لوحات فئة (خصوصي).", size=9.5, space_after=3)

    # UPDATED PAYMENT AND FEES CLAUSE
    add_p("آلية سداد الأقساط والرسوم:", size=10, bold=True, color=(41, 75, 98), space_after=1)
    add_p("نحيطكم علماً ونؤكد بأن سداد الأقساط الشهرية والتسهيلات التمويلية المتبقية لصالحكم الخاصة بهذه المركبات سيستمر كما هو قائم ومعمول به حالياً من قبل السيد / خميس هاشم الجبر دون أي تغيير أو تأخير في مواعيد استحقاقها المعتمدة، في حين تتحمل شركة جبريل لتوزيع الأغذية والمشروبات (ذ.م.م) كافة الرسوم الإدارية والمرورية المترتبة على نقل الملكية وتغيير اللوحات إلى خصوصي.", size=9.5, space_after=4)

    add_p("مرفق طيه بالملحق رقم (1) الكشف التفصيلي المتضمن أرقام اللوحات وأرقام الشاسيه (VIN) للـ (20) مركبة.", size=9.5, bold=True, color=(100, 116, 139), space_after=4)

    add_p("شاكرين ومقدرين لكم كريم تعاونكم وحسن تجاوبكم الدائم،،،", size=10, bold=True, space_after=8)

    # Signatures Table on Page 1
    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell_r = t_sig.cell(0, 0)
    p_r = cell_r.paragraphs[0]
    set_rtl(p_r, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    r_r1 = p_r.add_run("عن الطرف الأول (المتنازل والملتزم بالسداد):\n")
    r_r1.bold = True
    r_r1.font.color.rgb = RGBColor(41, 75, 98)
    r_r2 = p_r.add_run("شركة العراف لتأجير السيارات ذ.م.م\nالمدير المفوض: خميس هاشم الجبر\n\nالتوقيع والختم: .......................................")
    for r in (r_r1, r_r2):
        r.font.name = "Amiri"
        r.font.size = Pt(9.5)

    cell_l = t_sig.cell(0, 1)
    p_l = cell_l.paragraphs[0]
    set_rtl(p_l, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    r_l1 = p_l.add_run("عن الطرف الثاني (المتنازل إليه):\n")
    r_l1.bold = True
    r_l1.font.color.rgb = RGBColor(41, 75, 98)
    r_l2 = p_l.add_run("شركة جبريل لتوزيع الأغذية والمشروبات ذ.م.م\nالمدير المفوض: طارق خليفة العريبي\n\nالتوقيع والختم: .......................................")
    for r in (r_l1, r_l2):
        r.font.name = "Amiri"
        r.font.size = Pt(9.5)

    # PAGE BREAK FOR ANNEX
    doc.add_page_break()

    # PAGE 2: ANNEX 1
    p_an_meta = doc.add_paragraph()
    set_rtl(p_an_meta, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    r_an = p_an_meta.add_run("ملحق رقم (1) — تابع للكتاب رقم: ALARAF/ATT/2026-10")
    r_an.font.name = "Amiri"
    r_an.font.size = Pt(9)
    r_an.font.color.rgb = RGBColor(100, 116, 139)

    add_p("كشف حصر سيارات شانجان (20 مركبة) المطلوب نقل ملكيتها وتغيير لوحاتها لخصوصي", size=11, bold=True, color=(41, 75, 98), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    vehicles = [
        ("1", "10670", "LS5A2ASE6RD912602", "Alsvin", "2024"),
        ("2", "10852", "LS5A2ASE1RD910935", "Alsvin", "2024"),
        ("3", "10663", "LS5A2ASEXRD912568", "Alsvin", "2024"),
        ("4", "10854", "LS5A2ASE3RD910936", "Alsvin", "2024"),
        ("5", "10857", "LS5A2ASEXRD910934", "Alsvin", "2024"),
        ("6", "10853", "LS5A2ASE9RD910942", "Alsvin", "2024"),
        ("7", "10668", "LS5A2ASE7RD912608", "Alsvin", "2024"),
        ("8", "10850", "LS5A2ASE9RD910939", "Alsvin", "2024"),
        ("9", "10666", "LS5A2ASE8RD912570", "Alsvin", "2024"),
        ("10", "10851", "LS5A2ASE8RD910933", "Alsvin", "2024"),
        ("11", "10664", "LS5A2ASE1RD912605", "Alsvin", "2024"),
        ("12", "10672", "LS5A2ASEXRD912604", "Alsvin", "2024"),
        ("13", "10671", "LS5A2ASE1RD912569", "Alsvin", "2024"),
        ("14", "10858", "LS5A2ASE7RD910938", "Alsvin", "2024"),
        ("15", "10669", "LS5A2ASE6RD912597", "Alsvin", "2024"),
        ("16", "10667", "LS5A2ASE8RD912567", "Alsvin", "2024"),
        ("17", "10665", "LS5A2ASE0RD912563", "Alsvin", "2024"),
        ("18", "10856", "LS5A2ASE5RD910940", "Alsvin", "2024"),
        ("19", "10849", "LS5A2ASE7RD910941", "Alsvin", "2024"),
        ("20", "10855", "LS5A2ASE5RD910937", "Alsvin", "2024")
    ]

    t_veh = doc.add_table(rows=len(vehicles)+1, cols=5)
    t_veh.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_veh)

    headers = ["م", "رقم اللوحة", "رقم الشاسيه (VIN)", "الموديل", "سنة الصنع"]
    for col_idx, h in enumerate(headers):
        cell = t_veh.cell(0, col_idx)
        set_cell_background(cell, "294B62")
        set_cell_margins(cell, top=35, bottom=35, left=60, right=60)
        p = cell.paragraphs[0]
        set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        r = p.add_run(h)
        r.font.name = "Amiri"
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, data in enumerate(vehicles):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            cell = t_veh.cell(row_idx+1, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=25, bottom=25, left=50, right=50)
            p = cell.paragraphs[0]
            set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
            r = p.add_run(val)
            r.font.name = "Amiri"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(23, 50, 71)

    add_p("", space_after=4)

    # Signatures on Page 2
    t_sig2 = doc.add_table(rows=1, cols=2)
    t_sig2.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    c_r = t_sig2.cell(0, 0)
    p_r2 = c_r.paragraphs[0]
    set_rtl(p_r2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    r_r = p_r2.add_run("اعتماد المتنازل (شركة العراف لتأجير السيارات)\nالختم والتوقيع: .......................................")
    r_r.font.name = "Amiri"
    r_r.font.size = Pt(9)
    r_r.bold = True
    r_r.font.color.rgb = RGBColor(41, 75, 98)

    c_l = t_sig2.cell(0, 1)
    p_l2 = c_l.paragraphs[0]
    set_rtl(p_l2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    r_l = p_l2.add_run("اعتماد المتنازل إليه (شركة جبريل لتوزيع الأغذية)\nالختم والتوقيع: .......................................")
    r_l.font.name = "Amiri"
    r_l.font.size = Pt(9)
    r_l.bold = True
    r_l.font.color.rgb = RGBColor(41, 75, 98)

    out_docx = '/home/ubuntu/fleetifyapp/reports/letter_alattiya_changan_transfer.docx'
    doc.save(out_docx)
    print("Updated tight 2-page letter saved successfully!")

update_alattiya_letter()
