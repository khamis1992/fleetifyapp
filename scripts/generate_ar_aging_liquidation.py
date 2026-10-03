import os, csv
from datetime import datetime, date
from collections import defaultdict
from supabase import create_client
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

url = os.environ.get('SUPABASE_URL')
key = os.environ.get('SUPABASE_KEY')
supabase = create_client(url, key)
auth_res = supabase.auth.sign_in_with_password({'email': 'Khamis-1992@hotmail.com', 'password': '123456789'})
client = create_client(url, key)
client.postgrest.auth(auth_res.session.access_token)
company_id = '24bc0b21-4e2d-4413-9842-31719a3669f4'

cutoff = date(2026, 9, 30)

# Fetch customers
cust_res = client.table('customers').select('id, customer_code, first_name, last_name, company_name, phone, national_id').eq('company_id', company_id).execute()
customers = {}
for c in cust_res.data:
    name = c.get('company_name') or f"{c.get('first_name') or ''} {c.get('last_name') or ''}".strip() or 'عميل بدون اسم'
    customers[c['id']] = {
        'code': c.get('customer_code') or '',
        'name': name,
        'phone': c.get('phone') or '',
        'national_id': c.get('national_id') or ''
    }

# Fetch unpaid/partially paid active invoices
invoices = []
offset = 0
while True:
    res = client.table('invoices').select('id, customer_id, invoice_number, invoice_date, due_date, total_amount, paid_amount, balance_due, status').eq('company_id', company_id).not_.in_('status', ['cancelled', 'draft']).range(offset, offset+999).execute()
    invoices.extend(res.data)
    if len(res.data) < 1000:
        break
    offset += 1000

print(f"Total active invoices fetched: {len(invoices)}")

customer_aging = defaultdict(lambda: {
    'code': '', 'name': '', 'phone': '', 'national_id': '',
    'current_0_30': 0.0,
    'days_31_60': 0.0,
    'days_61_90': 0.0,
    'days_91_180': 0.0,
    'days_181_365': 0.0,
    'days_over_365': 0.0,
    'total_due': 0.0
})

for inv in invoices:
    bal = float(inv.get('balance_due') or 0)
    if bal <= 0.01:
        continue
    cid = inv.get('customer_id')
    cust = customers.get(cid, {'code': '', 'name': 'عميل غير مسجل', 'phone': '', 'national_id': ''})
    c_entry = customer_aging[cid]
    c_entry['code'] = cust['code']
    c_entry['name'] = cust['name']
    c_entry['phone'] = cust['phone']
    c_entry['national_id'] = cust['national_id']
    
    d_str = inv.get('due_date') or inv.get('invoice_date')
    if d_str:
        inv_date = datetime.strptime(str(d_str)[:10], '%Y-%m-%d').date()
    else:
        inv_date = date(2025, 1, 1)
    
    days_past = (cutoff - inv_date).days
    c_entry['total_due'] += bal
    if days_past <= 30:
        c_entry['current_0_30'] += bal
    elif days_past <= 60:
        c_entry['days_31_60'] += bal
    elif days_past <= 90:
        c_entry['days_61_90'] += bal
    elif days_past <= 180:
        c_entry['days_91_180'] += bal
    elif days_past <= 365:
        c_entry['days_181_365'] += bal
    else:
        c_entry['days_over_365'] += bal

# Recovery rates:
rates = {
    '0_30': 0.90,
    '31_60': 0.75,
    '61_90': 0.50,
    '91_180': 0.30,
    '181_365': 0.15,
    'over_365': 0.05
}

# Sort customers by total due descending
sorted_customers = sorted(customer_aging.values(), key=lambda x: x['total_due'], reverse=True)

csv_path = '/home/ubuntu/fleetifyapp/reports/accounts_receivable_aging_liquidation.csv'
xlsx_path = '/home/ubuntu/fleetifyapp/reports/accounts_receivable_aging_liquidation.xlsx'

with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.writer(f)
    writer.writerow(['م', 'كود العميل', 'اسم العميل', 'رقم الهاتف', 'الرقم القومي / الهوية', '1-30 يوم (90%)', '31-60 يوم (75%)', '61-90 يوم (50%)', '91-180 يوم (30%)', '181-365 يوم (15%)', 'أكثر من سنة (5%)', 'إجمالي الرصيد القائم', 'التحصيل التقديري المتوقع للتفليسة'])
    for idx, c in enumerate(sorted_customers, 1):
        est_rec = (
            c['current_0_30'] * rates['0_30'] +
            c['days_31_60'] * rates['31_60'] +
            c['days_61_90'] * rates['61_90'] +
            c['days_91_180'] * rates['91_180'] +
            c['days_181_365'] * rates['181_365'] +
            c['days_over_365'] * rates['over_365']
        )
        writer.writerow([
            idx, c['code'], c['name'], c['phone'], c['national_id'],
            f"{c['current_0_30']:,.2f}",
            f"{c['days_31_60']:,.2f}",
            f"{c['days_61_90']:,.2f}",
            f"{c['days_91_180']:,.2f}",
            f"{c['days_181_365']:,.2f}",
            f"{c['days_over_365']:,.2f}",
            f"{c['total_due']:,.2f}",
            f"{est_rec:,.2f}"
        ])

# Create XLSX
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'أعمار الذمم والتحصيل المتوقع'
ws.views.sheetView[0].rightToLeft = True

header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

with open(csv_path, 'r', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for r_idx, row in enumerate(reader, 1):
        ws.append(row)
        for c_idx, cell_value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if r_idx == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            else:
                cell.font = Font(name='Calibri', size=10)
                if c_idx in [1, 2, 4, 5]:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif c_idx >= 6:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='right', vertical='center')

for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(xlsx_path)

tot_0_30 = sum(c['current_0_30'] for c in sorted_customers)
tot_31_60 = sum(c['days_31_60'] for c in sorted_customers)
tot_61_90 = sum(c['days_61_90'] for c in sorted_customers)
tot_91_180 = sum(c['days_91_180'] for c in sorted_customers)
tot_181_365 = sum(c['days_181_365'] for c in sorted_customers)
tot_over_365 = sum(c['days_over_365'] for c in sorted_customers)
tot_due = sum(c['total_due'] for c in sorted_customers)
tot_rec = (
    tot_0_30 * rates['0_30'] +
    tot_31_60 * rates['31_60'] +
    tot_61_90 * rates['61_90'] +
    tot_91_180 * rates['91_180'] +
    tot_181_365 * rates['181_365'] +
    tot_over_365 * rates['over_365']
)

print(f"Total Outstanding: {tot_due:,.2f} QAR")
print(f"0-30 days: {tot_0_30:,.2f} QAR")
print(f"31-60 days: {tot_31_60:,.2f} QAR")
print(f"61-90 days: {tot_61_90:,.2f} QAR")
print(f"91-180 days: {tot_91_180:,.2f} QAR")
print(f"181-365 days: {tot_181_365:,.2f} QAR")
print(f"Over 365 days: {tot_over_365:,.2f} QAR")
print(f"Expected Recovery: {tot_rec:,.2f} QAR ({(tot_rec/tot_due)*100:.1f}%)")
